#!/usr/bin/env python3
"""Reproducible official food-data -> indexed, versioned SQLite packs.

Raw upstream downloads stay in .qa/country-food-sources. No keys, service calls
at runtime, diary information, or publication are involved in this build.
Requires Python 3.11+ and openpyxl for the two upstream Excel files.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
import gzip
import hashlib
import io
import json
import math
import re
import sqlite3
import time
import unicodedata
import urllib.request
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / '.qa/country-food-sources'
OUT = ROOT / 'assets/food-packs'
CHUNK_SIZE = 1024 * 1024
CNF_URL = 'https://open.canada.ca/data/dataset/1b6139bd-ed7e-4043-bc28-ff00e10f3109/resource/019f2a90-e3a9-489d-b6e1-f74f4ba1d006/download/cnf_fcen_all-files-data_2026.zip'
FINELI_URL = 'https://fineli.fi/fineli/content/file/49'
BLS_URL = 'https://blsdb.de/assets/uploads/BLS_4_0_2025_DE.zip'
USDA_URL = 'https://fdc.nal.usda.gov/fdc-datasets/FoodData_Central_branded_food_json_2026-04-30.zip'
CIQUAL_API = 'https://entrepot.recherche.data.gouv.fr/api/datasets/:persistentId/?persistentId=doi:10.57745/RDMHWY'


def fetch(url: str, path: Path) -> Path:
    """Atomic cache download. Do not use a partially received source."""
    if path.exists():
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    pending = path.with_suffix(path.suffix + '.partial')
    request = urllib.request.Request(url, headers={'User-Agent': 'T1Arc-food-pack-build/1.0'})
    with urllib.request.urlopen(request, timeout=90) as response, pending.open('wb') as out:
        while block := response.read(CHUNK_SIZE):
            out.write(block)
    pending.replace(path)
    print(f'Downloaded {path.name}: {path.stat().st_size:,} bytes', flush=True)
    return path


def sources(include_usda: bool = False, include_fineli: bool = False) -> dict[str, Path]:
    from concurrent.futures import ThreadPoolExecutor
    def ciqual():
        metadata = fetch(CIQUAL_API, CACHE / 'ciqual-2025-metadata.json')
        files = json.loads(metadata.read_text(encoding='utf8'))['data']['latestVersion']['files']
        item = next(f['dataFile'] for f in files if f['dataFile']['filename'].endswith('.xlsx') and '_FR_' in f['dataFile']['filename'])
        return fetch(f"https://entrepot.recherche.data.gouv.fr/api/access/datafile/{item['id']}", CACHE / 'ciqual-2025.xlsx')
    jobs = {
        'cnf': lambda: fetch(CNF_URL, CACHE / 'cnf-2026.zip'),
        'ciqual': ciqual,
        'bls': lambda: fetch(BLS_URL, CACHE / 'bls-4.0.zip'),
    }
    if include_fineli:
        jobs['fineli'] = lambda: fetch(FINELI_URL, CACHE / 'fineli-74.zip')
    if include_usda:
        jobs['usda-branded'] = lambda: fetch(USDA_URL, CACHE / 'usda-branded-2026-04.zip')
    with ThreadPoolExecutor(max_workers=5) as pool:
        running = {key: pool.submit(action) for key, action in jobs.items()}
        return {key: task.result() for key, task in running.items()}


def inspect(paths):
    import openpyxl
    for key, path in paths.items():
        print(key, path.name)
        if path.suffix == '.xlsx':
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            for sheet in book:
                print(sheet.title, sheet.max_row, sheet.max_column)
                for row in list(sheet.values)[:5]:
                    print(row)
        else:
            with zipfile.ZipFile(path) as archive:
                print(archive.namelist()[:35])
                for name in archive.namelist():
                    if name.lower().endswith(('.csv','.txt')):
                        with archive.open(name) as stream:
                            print(name, stream.read(1500).decode('utf-8-sig', errors='replace')[:900])
                    elif name.lower().endswith('.xlsx'):
                        book = openpyxl.load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=True)
                        for sheet in book:
                            print(name, sheet.title, sheet.max_row, sheet.max_column)
                            for row in list(sheet.values)[:4]:
                                print(row)


SPECS = {
    'cnf': dict(id='ca-cnf', provider='cnf', countryCode='CA', version='2026',
        label='Canadian Nutrient File 2026', sourceUrl=CNF_URL,
        licence='Open Government Licence - Canada', licenceUrl='https://open.canada.ca/en/open-government-licence-canada',
        attribution='Contains information licensed under the Open Government Licence – Canada. Canadian Nutrient File, Health Canada, 2026.',
        carbohydrateDefinition='by-difference', carbohydrateNote='Total carbohydrate by difference, including dietary fibre (CNF nutrient 205).', bundled=True),
    'ciqual': dict(id='fr-ciqual', provider='ciqual', countryCode='FR', version='2025',
        label='Anses Ciqual 2025', sourceUrl='https://doi.org/10.57745/RDMHWY',
        licence='Etalab Open Licence 2.0', licenceUrl='https://www.etalab.gouv.fr/licence-ouverte-open-licence/',
        attribution='Anses. 2025. Table de composition nutritionnelle des aliments Ciqual. Processed subset; not an official Anses product.',
        carbohydrateDefinition='available', carbohydrateNote='Available carbohydrate, as reported by Ciqual; no subtraction of fibre.', bundled=True),
    'bls': dict(id='de-bls', provider='bls', countryCode='DE', version='4.0-2025',
        label='Bundeslebensmittelschlüssel 4.0', sourceUrl=BLS_URL,
        licence='CC BY 4.0', licenceUrl='https://creativecommons.org/licenses/by/4.0/',
        attribution='Max Rubner-Institut (2025): Bundeslebensmittelschlüssel (BLS), Version 4.0 – Deutsche Nährstoffdatenbank. Karlsruhe. DOI: 10.25826/Data20251217-134202-0. Processed subset.',
        carbohydrateDefinition='available', carbohydrateNote='Available carbohydrate, including polyols (BLS CHO); not total carbohydrate including fibre.', bundled=True),
    'usda-branded': dict(id='us-usda-branded', provider='usda-fdc', countryCode='US', version='2026-04-30',
        label='USDA branded foods — April 2026', sourceUrl=USDA_URL,
        licence='CC0 1.0', licenceUrl='https://creativecommons.org/publicdomain/zero/1.0/',
        attribution='US Department of Agriculture, Agricultural Research Service. FoodData Central, branded foods April 2026. Processed offline snapshot; formulations may change.',
        carbohydrateDefinition='total', carbohydrateNote='Label total carbohydrate (USDA nutrient 1005), including fibre. Do not subtract fibre automatically.', bundled=True, optional=True, assetCompression='gzip', dataRevision=2),
}
KEYS = ('carbs','kcal','protein','fat','fibre','sugars','saturated_fat')


def tidy(value):
    return re.sub(r'\s+', ' ', '' if value is None else str(value)).strip()


def numeric(value, maximum=100.1):
    """Never turn missing/trace/below-LOQ source cells into a guessed zero."""
    if isinstance(value, bool):
        return None, 'missing'
    text = tidy(value).replace(',', '.')
    if text.lower() in ('traces','trace') or text.startswith('<'):
        return None, 'trace'
    try:
        number = float(text)
    except (ValueError, TypeError):
        return None, 'missing'
    if not math.isfinite(number) or number < 0 or number > maximum:
        return None, 'missing'
    return round(number, 6), 'reported'


def food(code, names, values, aliases=(), **extra):
    names = {lang: tidy(name)[:500] for lang, name in names.items() if tidy(name)}
    if not names or not tidy(code):
        return None
    parsed = [numeric(value, 1100 if key == 'kcal' else (200 if extra.get('basis_unit') == 'ml' else 100.1)) for key, value in zip(KEYS, values)]
    if all(value is None for value, _ in parsed):
        return None
    return dict(source_id=str(code), names=names, aliases=list(dict.fromkeys(tidy(v)[:300] for v in aliases if tidy(v)))[:8],
        values=[v for v, _ in parsed], quality=[q for _, q in parsed], **extra)


def csv_rows(archive, filename):
    raw = archive.read(filename)
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('cp1252')
    return csv.DictReader(io.StringIO(text, newline=None))


def cnf_rows(path):
    codes = dict(zip(('205','208','203','204','291','269','606'), range(7)))
    with zipfile.ZipFile(path) as archive:
        nutrients = {}
        for item in csv_rows(archive, 'Nutrient_Amount.csv'):
            index = codes.get(item['Nutrient_Code'])
            if index is not None:
                nutrients.setdefault(item['Food_Code'], [None]*7)[index] = item['Nutrient_Amount']
        measures = {r['Measure_Code']:r for r in csv_rows(archive,'Measure_Name.csv')}
        portions = {}
        for row in csv_rows(archive,'Measure_Weight_Conversion.csv'):
            amount, _ = numeric(row['Measure_Weight_Conversion'], 5000)
            if row['Measure_Type_Code'] == '6' and amount and row['Food_Code'] not in portions:
                label = measures.get(row['Measure_Code'],{}).get('Measure_Description_and_Unit_EN')
                if label:
                    portions[row['Food_Code']] = dict(serving_amount=amount, serving_unit='g', serving_label=tidy(label))
        for row in csv_rows(archive, 'Food_Name.csv'):
            yield food(row['Food_Code'], {'en':row['Food_Description_EN'],'fr':row['Food_Description_FR']},
                nutrients.get(row['Food_Code'],[None]*7),
                [row['Alternate_Description_EN'],row['Alternate_Description_FR']],
                **portions.get(row['Food_Code'],{}))


def ciqual_rows(path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = iter(book.worksheets[0].values)
    header = [tidy(v) for v in next(rows)]
    expected = {6:'alim_code',7:'alim_nom_fr',10:'Energie,',14:'Protéines,',16:'Glucides',17:'Lipides',18:'Sucres',26:'Fibres',31:'AG saturés'}
    for index, prefix in expected.items():
        if not header[index].startswith(prefix):
            raise ValueError(f'Ciqual schema changed: column {index}: {header[index]}')
    for row in rows:
        yield food(row[6], {'fr':row[7]}, [row[i] for i in [16,10,14,17,26,18,31]], [row[4],row[5]])
    book.close()


def bls_search_aliases(code, german, english):
    """Expose bread compounds to token-prefix search, without aliasing breadfruit.

    Restrict to source B-coded foods and explicit source-language bread names;
    this does not translate arbitrary foods or infer a nutritional category.
    """
    if tidy(code).startswith('B') and (
        re.search(r'\b[\w-]*brot\b', tidy(german), re.IGNORECASE)
        or re.search(r'\bbread\b', tidy(english), re.IGNORECASE)
    ):
        return ['Brot', 'bread']
    return []


def bls_rows(path):
    import openpyxl
    with zipfile.ZipFile(path) as archive:
        name = next(n for n in archive.namelist() if 'Daten_' in n and n.endswith('.xlsx'))
        book = openpyxl.load_workbook(io.BytesIO(archive.read(name)),read_only=True,data_only=True)
        rows = iter(book.active.values)
        header = next(rows)
        positions = [next(i for i,v in enumerate(header) if str(v).startswith(code+' ') and '/100g]' in str(v)) for code in ['CHO','ENERCC','PROT625','FAT','FIBT','SUGAR','FASAT']]
        for row in rows:
            yield food(row[0], {'de':row[1],'en':row[2]}, [row[i] for i in positions],
                bls_search_aliases(row[0], row[1], row[2]))
        book.close()


def stream_json_foods(stream):
    """Stream the USDA top-level food array without a multi-GB JSON allocation."""
    decoder = json.JSONDecoder()
    buffer = ''
    while '[' not in buffer:
        block = stream.read(65536)
        if not block:
            raise ValueError('No USDA food array')
        buffer += block
    position = buffer.index('[')+1
    while True:
        while position < len(buffer) and buffer[position] in ' \r\n\t,':
            position += 1
        if position < len(buffer) and buffer[position] == ']':
            return
        try:
            value, position = decoder.raw_decode(buffer, position)
            yield value
        except json.JSONDecodeError:
            buffer = buffer[position:]
            position = 0
            block = stream.read(CHUNK_SIZE)
            if not block:
                raise ValueError('Truncated USDA JSON')
            buffer += block
            if len(buffer) > 16*CHUNK_SIZE:
                raise ValueError('Unexpectedly large USDA record')


USDA_NUTRIENT_IDS = {1003, 1004, 1005, 1008, 1062, 1079, 2000, 1258, 2047, 2048}
USDA_IMPORT_STATS = {}


def usda_nutrients(items):
    """Only equal same-ID observations may collapse; list order is not authority."""
    groups = {}
    for item in items:
        nutrient_id = item.get('nutrient', {}).get('id')
        if nutrient_id not in USDA_NUTRIENT_IDS:
            continue
        value = item.get('amount')
        if value is None:
            continue
        # A less-than source is a bound, not a measured exact amount.
        if item.get('foodNutrientDerivation', {}).get('code') in ('LCGL', 'LCSL'):
            value = f'<{value}'
        groups.setdefault(nutrient_id, []).append(value)
    values, conflicts = {}, set()
    for nutrient_id, observations in groups.items():
        signatures = set()
        for value in observations:
            try:
                number = float(value)
                signature = ('number', number) if math.isfinite(number) and not isinstance(value, bool) else ('invalid', str(value))
            except (TypeError, ValueError):
                signature = ('source', str(value))
            signatures.add(signature)
        if len(signatures) > 1:
            values[nutrient_id] = None
            conflicts.add(nutrient_id)
        else:
            values[nutrient_id] = observations[0]
    return values, conflicts


def usda_food(row):
    values, conflicts = usda_nutrients(row.get('foodNutrients', []))
    # Branded foodNutrients are standardized per 100 of the declared source unit.
    unit = {'g': 'g', 'grm': 'g', 'ml': 'ml', 'mlt': 'ml'}.get(str(row.get('servingSizeUnit', '')).lower())
    if not unit:
        return None
    # Distinct energy methods are not duplicate observations. Never fall through
    # a conflicting selected ID to make that conflict disappear.
    energy_id = next((key for key in (1008, 2048, 2047, 1062) if key in values), None)
    kcal = values.get(energy_id)
    if energy_id == 1062:
        parsed, quality = numeric(kcal, 4602.4)
        kcal = parsed / 4.184 if parsed is not None else ('<unknown' if quality == 'trace' else None)
    nutrient_ids = (1005, energy_id, 1003, 1004, 1079, 2000, 1258)
    conflict_fields = [key for key, nutrient_id in zip(KEYS, nutrient_ids) if nutrient_id in conflicts]
    amount, _ = numeric(row.get('servingSize'), 5000)
    portion = dict(serving_amount=amount, serving_unit=unit, serving_label=tidy(row.get('householdServingFullText'))[:300]) if amount else {}
    barcode = tidy(row.get('gtinUpc'))
    barcode = barcode if len(barcode) == 8 else barcode.zfill(14)
    return food(row.get('fdcId'), {'en': row.get('description')},
        [values.get(1005), kcal, values.get(1003), values.get(1004), values.get(1079), values.get(2000), values.get(1258)],
        [row.get('brandedFoodCategory')], brand=tidy(row.get('brandName') or row.get('brandOwner'))[:300],
        barcode=barcode, source_date=usda_date(row.get('publicationDate')), basis_unit=unit,
        nutrient_conflicts=conflict_fields,
        energy_definition='atwater-specific' if energy_id == 2048 else ('atwater-general' if energy_id == 2047 else 'reported'), **portion)


def usda_rows(path):
    USDA_IMPORT_STATS.clear()
    USDA_IMPORT_STATS.update(conflictingSourceRecords=0, unsupportedBasisRecords=0)
    with zipfile.ZipFile(path) as archive:
        name = next(n for n in archive.namelist() if n.endswith('.json'))
        with archive.open(name) as source:
            for row in stream_json_foods(io.TextIOWrapper(source, encoding='utf8')):
                if row.get('marketCountry') != 'United States':
                    continue
                barcode = tidy(row.get('gtinUpc'))
                if not re.fullmatch(r'(?:\d{8}|\d{12,14})',barcode):
                    continue
                if usda_nutrients(row.get('foodNutrients', []))[1]:
                    USDA_IMPORT_STATS['conflictingSourceRecords'] += 1
                if str(row.get('servingSizeUnit', '')).lower() not in ('g', 'grm', 'ml', 'mlt'):
                    USDA_IMPORT_STATS['unsupportedBasisRecords'] += 1
                yield usda_food(row)


def extract_usda_test_fixtures(path):
    targets, foods = {2517161, 2738631}, []
    with zipfile.ZipFile(path) as archive:
        with archive.open(next(name for name in archive.namelist() if name.endswith('.json'))) as source:
            for row in stream_json_foods(io.TextIOWrapper(source, encoding='utf8')):
                if row.get('fdcId') in targets:
                    foods.append(row)
                    targets.remove(row['fdcId'])
                    if not targets:
                        break
    if targets:
        raise ValueError(f'Missing source fixtures: {targets}')
    target = ROOT / 'tests/fixtures/usda-branded-cheerios-2026-04.json'
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(dict(sourceUrl=USDA_URL, sourceSha256=hash_file(path)[0], foods=foods), ensure_ascii=False, indent=2)+'\n', encoding='utf8')
    print(f'Extracted exact source fixtures: {target.name}', flush=True)


def usda_date(value):
    text = tidy(value)
    for pattern in ('%m/%d/%Y','%Y-%m-%d'):
        try:
            return datetime.strptime(text,pattern).date().isoformat()
        except ValueError:
            pass
    return ''


def search_text(values):
    text = unicodedata.normalize('NFKD',' '.join(values)).lower()
    return re.sub(r'[^\w]+',' ',''.join(c for c in text if not unicodedata.combining(c)),flags=re.UNICODE).strip()


def encode(value):
    return json.dumps(value,ensure_ascii=False,separators=(',',':'),sort_keys=True)


def hash_file(path):
    digest = hashlib.sha256()
    chunks = []
    with path.open('rb') as stream:
        while block := stream.read(CHUNK_SIZE):
            digest.update(block)
            chunks.append(hashlib.sha256(block).hexdigest())
    return digest.hexdigest(), chunks


def build(key, path, accept_source_change=False):
    spec = SPECS[key]
    source_hash = hash_file(path)[0]
    trusted_manifest = ROOT / 'src/data/food/country-packs.manifest.json'
    if trusted_manifest.exists() and not accept_source_change:
        previous = next((p for p in json.loads(trusted_manifest.read_text(encoding='utf8')) if p['id']==spec['id']),None)
        if previous and previous['version']==spec['version'] and previous['sourceSha256'] != source_hash:
            raise ValueError(f'{key}: source bytes changed for the pinned version. Review upstream change before --accept-source-change.')
    destination = ROOT / '.qa/country-food-release-assets' if spec.get('optional') else OUT
    destination.mkdir(parents=True, exist_ok=True)
    target = destination / f"{spec['id']}-{spec['version']}.db"
    pending = target.with_suffix('.building.db')
    # This is only our generated temporary file, never an input/user database.
    if pending.exists():
        pending.unlink()
    db = sqlite3.connect(pending)
    db.executescript('''PRAGMA page_size=4096; PRAGMA journal_mode=DELETE; PRAGMA user_version=1;
      CREATE TABLE pack_metadata (key TEXT PRIMARY KEY,value TEXT NOT NULL) WITHOUT ROWID;
      CREATE TABLE foods (source_id TEXT PRIMARY KEY, names_json TEXT NOT NULL, aliases_json TEXT NOT NULL,
        brand TEXT,barcode TEXT,source_date TEXT,carbs REAL,kcal REAL,protein REAL,fat REAL,fibre REAL,sugars REAL,saturated_fat REAL,
        quality_json TEXT NOT NULL,serving_amount REAL,serving_unit TEXT,serving_label TEXT,energy_definition TEXT NOT NULL,
        basis_unit TEXT NOT NULL DEFAULT 'g',nutrient_conflicts_json TEXT NOT NULL DEFAULT '[]',
        CHECK(basis_unit IN ('g','ml')), CHECK(json_valid(nutrient_conflicts_json)),
        CHECK(carbs IS NULL OR carbs BETWEEN 0 AND CASE basis_unit WHEN 'ml' THEN 200 ELSE 100.1 END),CHECK(kcal IS NULL OR kcal BETWEEN 0 AND 1100),
        CHECK(protein IS NULL OR protein BETWEEN 0 AND CASE basis_unit WHEN 'ml' THEN 200 ELSE 100.1 END),CHECK(fat IS NULL OR fat BETWEEN 0 AND CASE basis_unit WHEN 'ml' THEN 200 ELSE 100.1 END),
        CHECK(fibre IS NULL OR fibre BETWEEN 0 AND CASE basis_unit WHEN 'ml' THEN 200 ELSE 100.1 END),CHECK(sugars IS NULL OR sugars BETWEEN 0 AND CASE basis_unit WHEN 'ml' THEN 200 ELSE 100.1 END),
        CHECK(saturated_fat IS NULL OR saturated_fat BETWEEN 0 AND CASE basis_unit WHEN 'ml' THEN 200 ELSE 100.1 END),
        CHECK(serving_amount IS NULL OR (serving_amount>0 AND serving_amount<=5000 AND serving_unit IN ('g','ml'))));
      CREATE UNIQUE INDEX foods_barcode ON foods(barcode) WHERE barcode IS NOT NULL;
    ''')
    started = time.perf_counter()
    received = rejected = duplicates = 0
    iterator = {'cnf':cnf_rows,'ciqual':ciqual_rows,'bls':bls_rows,'usda-branded':usda_rows}[key]
    for row in iterator(path):
        received += 1
        if row is None:
            rejected += 1
            continue
        values = (row['source_id'],encode(row['names']),encode(row['aliases']),row.get('brand'),row.get('barcode'),row.get('source_date'),
            *row['values'],encode(row['quality']),row.get('serving_amount'),row.get('serving_unit'),row.get('serving_label'),row.get('energy_definition','reported'),
            row.get('basis_unit','g'),encode(row.get('nutrient_conflicts',[])))
        if row.get('barcode'):
            existing = db.execute('SELECT source_id,source_date FROM foods WHERE barcode=?',(row['barcode'],)).fetchone()
            if existing:
                duplicates += 1
                if (row.get('source_date',''),int(row['source_id'])) <= (existing[1] or '',int(existing[0])):
                    continue
                db.execute('DELETE FROM foods WHERE source_id=?',(existing[0],))
        db.execute('INSERT INTO foods VALUES ('+','.join('?' for _ in values)+')',values)
        if received % 50000 == 0:
            db.commit()
            print(f'{key}: {received:,} records processed',flush=True)
    db.commit()
    db.execute("CREATE VIRTUAL TABLE food_search USING fts5(text,content='',tokenize='unicode61 remove_diacritics 2')")
    for rowid,names,aliases,brand in db.execute('SELECT rowid,names_json,aliases_json,brand FROM foods ORDER BY source_id'):
        body = search_text([*json.loads(names).values(),*json.loads(aliases),brand or ''])
        db.execute('INSERT INTO food_search(rowid,text) VALUES (?,?)',(rowid,body))
    count = db.execute('SELECT count(*) FROM foods').fetchone()[0]
    metadata = {**spec,'schemaVersion':1,'foodCount':count,'sourceSha256':source_hash,
        'sourceFile':path.name,'basisAmount':100,'basisUnit':'per-row' if key == 'usda-branded' else 'g','transformation':'T1 Arc selected nutrients; missing/trace/LOQ and contradictory same-nutrient observations kept non-numeric; no inferred nutrition or density. Source g/GRM and ml/MLT units normalized. Names normalized for whitespace; search removes accents. Stable source IDs retained.'}
    if key == 'usda-branded':
        metadata.update(USDA_IMPORT_STATS)
        metadata['retainedConflictFoods'] = db.execute("SELECT count(*) FROM foods WHERE nutrient_conflicts_json != '[]'").fetchone()[0]
    db.execute('INSERT INTO pack_metadata VALUES (?,?)',('manifest',encode(metadata)))
    db.commit()
    db.execute('VACUUM')
    assert db.execute('PRAGMA quick_check').fetchone()[0] == 'ok'
    samples = ['bread','rice','milk','pain','pomme','brot','hafer','chicken']
    times = []
    for query in samples*20:
        start = time.perf_counter()
        db.execute('SELECT f.source_id FROM food_search s JOIN foods f ON f.rowid=s.rowid WHERE food_search MATCH ? ORDER BY bm25(food_search) LIMIT 20',(query+'*',)).fetchall()
        times.append((time.perf_counter()-start)*1000)
    barcode_count = db.execute('SELECT count(*) FROM foods WHERE barcode IS NOT NULL').fetchone()[0]
    db.close()
    pending.replace(target)
    digest,chunks = hash_file(target)
    compressed = target.with_suffix('.db.gz')
    with target.open('rb') as source, compressed.open('wb') as dest:
        with gzip.GzipFile(fileobj=dest,mode='wb',mtime=0,filename='') as zipped:
            while block := source.read(CHUNK_SIZE):
                zipped.write(block)
    entry = dict(**metadata,fileName=target.name,sizeBytes=target.stat().st_size,sha256=digest,chunkSize=CHUNK_SIZE,chunkSha256=chunks,
        compressedBytes=compressed.stat().st_size)
    if spec.get('assetCompression') == 'gzip':
        import shutil
        OUT.mkdir(parents=True,exist_ok=True)
        shutil.copyfile(compressed, OUT / compressed.name)
        compressed_hash, compressed_chunks = hash_file(compressed)
        entry.update(assetFileName=compressed.name, compressedSha256=compressed_hash, compressedChunkSha256=compressed_chunks)
    report = dict(id=spec['id'],received=received,rejected=rejected,duplicateBarcodes=duplicates,foodCount=count,barcodeCount=barcode_count,
        installedBytes=target.stat().st_size,gzipBytes=compressed.stat().st_size,buildSeconds=round(time.perf_counter()-started,2),
        queryMedianMs=round(sorted(times)[len(times)//2],3),queryP95Ms=round(sorted(times)[int(len(times)*.95)],3),
        benchmarkEnvironment='Desktop Python SQLite; warm queries, not Android memory/latency evidence',sha256=digest,sqliteVersion=sqlite3.sqlite_version)
    if key == 'usda-branded':
        report.update(USDA_IMPORT_STATS, retainedConflictFoods=metadata['retainedConflictFoods'])
    (destination / (spec['id']+'.benchmark.json')).write_text(json.dumps(report,indent=2)+'\n',encoding='utf8')
    (destination / (spec['id']+'.manifest.json')).write_text(json.dumps(entry,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps(report),flush=True)
    # Compress only for release benchmarking. APK assets are native SQLite files.
    if spec['bundled'] and not spec.get('optional'):
        compressed.unlink()
    return entry


def publish_local_manifest(entries):
    """Generate only local source artifacts. Never upload a release."""
    manifest_path = ROOT / 'src/data/food/country-packs.manifest.json'
    previous = json.loads(manifest_path.read_text(encoding='utf8')) if manifest_path.exists() else []
    merged = {entry['id']:entry for entry in [*previous,*entries]}
    manifest_path.write_text(json.dumps(list(merged.values()),ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    lines = ['// Generated by scripts/build_country_food_packs.py. Metro .db assets remain outside JS memory.',
        'export const COUNTRY_PACK_ASSETS: Record<string, () => number> = {']
    for entry in merged.values():
        if entry['bundled']:
            if entry.get('assetCompression'):
                lines.append('  // eslint-disable-next-line @typescript-eslint/no-require-imports -- Metro resolves this audited compressed native asset.')
            lines.append(f"  '{entry['id']}': () => require('../../../assets/food-packs/{entry.get('assetFileName',entry['fileName'])}'),")
    lines.append('};\n')
    (ROOT / 'src/data/food/countryPackAssets.ts').write_text('\n'.join(lines),encoding='utf8')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inspect', action='store_true')
    parser.add_argument('--usda-branded', action='store_true')
    parser.add_argument('--fineli', action='store_true', help='Optional upstream download (server may refuse access)')
    parser.add_argument('--only', choices=['cnf','ciqual','bls','usda-branded'])
    parser.add_argument('--accept-source-change',action='store_true',help='Explicitly accept reviewed new source bytes for an existing version')
    parser.add_argument('--extract-usda-fixtures',action='store_true',help='Extract exact public source records used by regression tests')
    args = parser.parse_args()
    paths = sources(args.usda_branded,args.fineli)
    if args.extract_usda_fixtures:
        extract_usda_test_fixtures(paths['usda-branded'])
    elif args.inspect:
        inspect(paths)
    else:
        entries = [build(key,path,args.accept_source_change) for key,path in paths.items() if key in SPECS and (not args.only or key == args.only)]
        publish_local_manifest(entries)
